"""Spark readers for supported source file formats."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

import fsspec
from pyspark.sql import DataFrame, SparkSession
from pyspark.sql.types import StringType, StructField, StructType


class SourceReadError(RuntimeError):
    """A source file could not be read or parsed."""


def read_source(
    spark: SparkSession, path: str | Path, file_format: str, options: dict[str, Any]
) -> DataFrame:
    """Read one source file, keeping CSV and XLSX source cells as strings."""
    try:
        if file_format == "xlsx":
            return _read_xlsx(spark, str(path), options)
        if file_format == "csv":
            _validate_csv(str(path), options)
        reader = spark.read
        for key, value in options.items():
            reader = reader.option(key, str(value))
        if file_format == "csv":
            reader = reader.option("inferSchema", "false")
        if file_format in {"csv", "json"}:
            reader = reader.option("mode", "FAILFAST")
        spark_path = str(path)
        if spark_path.startswith("s3://"):
            spark_path = "s3a://" + spark_path[len("s3://") :]
        return reader.format(file_format).load(spark_path)
    except Exception as exc:
        raise SourceReadError(f"Failed to read source as {file_format}") from exc


def _read_xlsx(spark: SparkSession, path: str, options: dict[str, Any]) -> DataFrame:
    from openpyxl import load_workbook

    filesystem, inner_path = fsspec.core.url_to_fs(path)
    with filesystem.open(inner_path, "rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=False)
        try:
            sheet_option = options.get("sheet", 0)
            sheet = (
                workbook.worksheets[int(sheet_option)]
                if str(sheet_option).isdigit()
                else workbook[str(sheet_option)]
            )
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                raise ValueError("XLSX sheet is empty or has no header row")
            names = [
                str(value).strip() if value is not None else f"_c{index}"
                for index, value in enumerate(headers)
            ]
            if len(set(names)) != len(names):
                raise ValueError("XLSX header contains duplicate column names")
            records = [
                [None if value is None else str(value) for value in row] for row in rows
            ]
            schema = StructType(
                [StructField(name, StringType(), True) for name in names]
            )
            return spark.createDataFrame(records, schema=schema)
        finally:
            workbook.close()


def _validate_csv(path: str, options: dict[str, Any]) -> None:
    """Reject malformed quoting/row widths that Spark CSV may silently tolerate."""
    with fsspec.open(
        path, "rt", encoding=options.get("encoding", "utf-8"), newline=""
    ) as stream:
        rows = csv.reader(stream, delimiter=options.get("sep", ","), strict=True)
        header = next(rows, None)
        if not header or len(header) != len({name.lower() for name in header}):
            raise ValueError("Missing or duplicate CSV header")
        for row in rows:
            if len(row) != len(header):
                raise ValueError("CSV record width differs from header")
